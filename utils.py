import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
import io
import re
import unicodedata
from datetime import datetime

def remover_acentos(texto):
    if not texto: return ""
    return unicodedata.normalize('NFD', str(texto)).encode('ascii', 'ignore').decode('utf-8').lower()

def extrair_data_referencia(conteudo_bytes):
    try:
        texto = conteudo_bytes.decode('latin1', errors='ignore')
        match = re.search(r'Gerado em;?\s*(\d{2}/\d{2}/\d{4})', texto, re.IGNORECASE)
        if match:
            return datetime.strptime(match.group(1), '%d/%m/%Y')
    except Exception:
        pass
    return datetime.now()

def carregar_e_limpar_csv(file_bytes):
    data_ref = extrair_data_referencia(file_bytes)
    lines = file_bytes.decode('latin1', errors='ignore').splitlines()
    
    header_idx = 0
    for idx, line in enumerate(lines[:30]):
        if line.startswith("Nome;") or ";Nome;" in line or line.startswith('"Nome";'):
            header_idx = idx
            break
            
    df = pd.read_csv(
        io.BytesIO(file_bytes),
        sep=';',
        encoding='latin1',
        skiprows=header_idx,
        dtype=str
    )
    
    df.columns = [str(c).strip().replace('"', '') for c in df.columns]
    
    def clean_val(x):
        if pd.notnull(x):
            s = str(x).strip().replace('"', '')
            return s if s not in ['', '-', 'nan', 'NaN', 'None'] else None
        return None

    try:
        df = df.map(clean_val)
    except AttributeError:
        df = df.applymap(clean_val)

    return df, data_ref

def montar_endereco(row):
    rua = str(row.get('Rua', '') or '').strip()
    num = str(row.get('Número', '') or '').strip()
    comp = str(row.get('Complemento', '') or '').strip()
    partes = [p for p in [rua, num, comp] if p and p not in ['None', 'nan', '', '-']]
    return ", ".join(partes) if partes else "Endereço não informado"

def buscar_coluna_flexivel(df, termos):
    for col in df.columns:
        col_norm = remover_acentos(col)
        for t in termos:
            t_norm = remover_acentos(t)
            if t_norm in col_norm:
                return col
    return None

def extrair_datas_validas(visitas_txt, data_ref, janela_dias=365):
    if not visitas_txt: return []
    datas_v = re.findall(r'\d{2}/\d{2}/\d{4}', str(visitas_txt))
    dts = []
    for d in datas_v:
        try:
            dt = datetime.strptime(d, '%d/%m/%Y')
            if 0 <= (data_ref - dt).days <= janela_dias:
                dts.append(dt)
        except: pass
    return sorted(list(set(dts)))

def gerar_excel_padrao(df_micro, df_nominal, titulo, data_ref, cols_status, col_score):
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    
    blue_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    gray_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    border_thin = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                         top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))

    # Aba 1: Dashboard_Microarea
    ws1 = wb.active
    ws1.title = "Dashboard_Microarea"
    ws1.append([titulo])
    ws1.append([f"Data de Referência: {data_ref.strftime('%d/%m/%Y')} | Unidade de Saúde / e-SUS APS"])
    ws1.append([])
    
    ws1.cell(row=1, column=1).font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    ws1.cell(row=2, column=1).font = Font(name="Calibri", size=10, italic=True)

    headers_micro = list(df_micro.columns)
    ws1.append(headers_micro)
    
    for col_num in range(1, len(headers_micro) + 1):
        cell = ws1.cell(row=4, column=col_num)
        cell.fill = blue_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in dataframe_to_rows(df_micro, index=False, header=False):
        ws1.append(r)

    for row in ws1.iter_rows(min_row=5, max_row=ws1.max_row, min_col=2, max_col=len(headers_micro)):
        for cell in row:
            if cell.column >= 3:
                cell.number_format = '0.0"%"'
            else:
                cell.number_format = '#,##0'
            cell.border = border_thin

    # Aba 2: Auditoria_Nominal
    ws2 = wb.create_sheet(title="Auditoria_Nominal")
    headers_nom = list(df_nominal.columns)
    ws2.append(headers_nom)
    
    for col_num in range(1, len(headers_nom) + 1):
        cell = ws2.cell(row=1, column=col_num)
        cell.fill = blue_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, r in enumerate(dataframe_to_rows(df_nominal, index=False, header=False), start=2):
        ws2.append(r)
        for c_idx, col_name in enumerate(headers_nom, start=1):
            cell = ws2.cell(row=r_idx, column=c_idx)
            cell.border = border_thin
            val = str(cell.value or '')
            
            if col_name in cols_status or "Prática" in col_name or "status" in col_name:
                if "Atendida" in val:
                    cell.fill = green_fill
                elif "Não atendida" in val:
                    cell.fill = red_fill
                elif "Fora da faixa" in val or "Aguardando" in val or "Não se aplica" in val or "acompanhamento" in val:
                    cell.fill = gray_fill
            elif col_name == col_score:
                try:
                    cell.value = float(cell.value)
                    cell.number_format = '0.0"%"'
                except: pass

    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(output)
    return output.getvalue()
