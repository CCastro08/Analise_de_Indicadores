import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
import io
import re
from datetime import datetime

st.set_page_config(page_title="Auditoria e-SUS APS - SAPS/MS", layout="wide")

st.title("🏥 Sistema de Auditoria de Indicadores e-SUS APS (SAPS/MS)")
st.subheader("Processamento Automático e Determinístico (C2, C3, C4, C5, C6, C7)")

# --- UTILITÁRIOS DE LEITURA E HIGIENIZAÇÃO DE CSV ---

def extrair_data_referencia(conteudo_bytes):
    try:
        texto = conteudo_bytes.decode('latin1', errors='ignore')
        match = re.search(r'Gerado em;?\s*(\d{2}/\d{2}/\d{4})', texto, re.IGNORECASE)
        if match:
            return datetime.strptime(match.group(1), '%d/%m/%Y')
    except Exception:
        pass
    return datetime.now()

def carregar_e_limpar_csv(file_bytes):
    data_ref = extrair_data_referencia(file_bytes)
    lines = file_bytes.decode('latin1', errors='ignore').splitlines()
    
    header_idx = 0
    for idx, line in enumerate(lines[:30]):
        if line.startswith("Nome;") or ";Nome;" in line or line.startswith('"Nome";'):
            header_idx = idx
            break
            
    df = pd.read_csv(
        io.BytesIO(file_bytes),
        sep=';',
        encoding='latin1',
        skiprows=header_idx,
        dtype=str
    )
    
    df.columns = [str(c).strip().replace('"', '') for c in df.columns]
    df = df.map(lambda x: str(x).strip().replace('"', '') if pd.notnull(x) else x)
    df = df.replace({'': None, '-': None, 'nan': None, 'NaN': None, 'None': None})
    return df, data_ref

def montar_endereco(row):
    rua = str(row.get('Rua', '') or '').strip()
    num = str(row.get('Número', '') or '').strip()
    comp = str(row.get('Complemento', '') or '').strip()
    partes = [p for p in [rua, num, comp] if p and p not in ['None', 'nan', '', '-']]
    return ", ".join(partes) if partes else "Endereço não informado"

def extrair_datas_validas(visitas_txt, data_ref, janela_dias=365):
    if not visitas_txt: return []
    datas_v = re.findall(r'\d{2}/\d{2}/\d{4}', str(visitas_txt))
    dts = []
    for d in datas_v:
        try:
            dt = datetime.strptime(d, '%d/%m/%Y')
            if 0 <= (data_ref - dt).days <= janela_dias:
                dts.append(dt)
        except: pass
    return sorted(list(set(dts)))

def buscar_coluna_flexivel(df, termos):
    for col in df.columns:
        col_norm = col.lower().replace('á','a').replace('é','e').replace('í','i').replace('ó','o').replace('ú','u').replace('ã','a').replace('ç','c')
        for t in termos:
            if t in col_norm:
                return col
    return None

# --- GERADOR DE EXCEL PADRONIZADO ---

def gerar_excel_padrao(df_micro, df_nominal, titulo, data_ref, cols_status, col_score):
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    
    blue_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    gray_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    border_thin = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                         top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))

    # Aba 1: Dashboard_Microarea
    ws1 = wb.active
    ws1.title = "Dashboard_Microarea"
    ws1.append([titulo])
    ws1.append([f"Data de Referência: {data_ref.strftime('%d/%m/%Y')} | Unidade de Saúde / e-SUS APS"])
    ws1.append([])
    
    ws1.cell(row=1, column=1).font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    ws1.cell(row=2, column=1).font = Font(name="Calibri", size=10, italic=True)

    headers_micro = list(df_micro.columns)
    ws1.append(headers_micro)
    
    for col_num in range(1, len(headers_micro) + 1):
        cell = ws1.cell(row=4, column=col_num)
        cell.fill = blue_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in dataframe_to_rows(df_micro, index=False, header=False):
        ws1.append(r)

    for row in ws1.iter_rows(min_row=5, max_row=ws1.max_row, min_col=3, max_col=len(headers_micro)):
        for cell in row:
            cell.number_format = '0.0"%"'
            cell.border = border_thin

    # Aba 2: Auditoria_Nominal
    ws2 = wb.create_sheet(title="Auditoria_Nominal")
    headers_nom = list(df_nominal.columns)
    ws2.append(headers_nom)
    
    for col_num in range(1, len(headers_nom) + 1):
        cell = ws2.cell(row=1, column=col_num)
        cell.fill = blue_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, r in enumerate(dataframe_to_rows(df_nominal, index=False, header=False), start=2):
        ws2.append(r)
        for c_idx, col_name in enumerate(headers_nom, start=1):
            cell = ws2.cell(row=r_idx, column=c_idx)
            cell.border = border_thin
            val = str(cell.value or '')
            
            if col_name in cols_status or "Prática" in col_name or "status" in col_name:
                if "Atendida" in val:
                    cell.fill = green_fill
                elif "Não atendida" in val:
                    cell.fill = red_fill
                elif "Fora da faixa" in val or "Aguardando" in val or "Não se aplica" in val:
                    cell.fill = gray_fill
            elif col_name == col_score:
                try:
                    cell.value = float(cell.value)
                    cell.number_format = '0.0"%"'
                except: pass

    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(output)
    return output.getvalue()

# --- MÓDULO INDICADOR C3 (GESTANTES E PUÉRPERAS) ---

def processar_c3(df, data_ref):
    df['Endereço'] = df.apply(montar_endereco, axis=1)
    col_dum = buscar_coluna_flexivel(df, ['dum'])
    col_ig = buscar_coluna_flexivel(df, ['idade gestacional', 'ig atual', 'ig'])
    
    def extrair_ig_dias(row):
        if col_dum and pd.notnull(row.get(col_dum)):
            try:
                dt = datetime.strptime(str(row.get(col_dum))[:10], '%d/%m/%Y')
                dias = (data_ref - dt).days
                if 0 <= dias <= 336: return dias
            except: pass
        if col_ig and pd.notnull(row.get(col_ig)):
            txt = str(row.get(col_ig)).lower()
            m_sem = re.search(r'(\d+)\s*s', txt)
            if m_sem:
                semanas = int(m_sem.group(1))
                m_dias = re.search(r'(\d+)\s*d', txt)
                dias_ext = int(m_dias.group(1)) if m_dias else 0
                return (semanas * 7) + dias_ext
        return 140 # Valor padrão (20 semanas) caso ausente

    df['IG_Dias_Calc'] = df.apply(extrair_ig_dias, axis=1)
    df['IG_Semanas'] = df['IG_Dias_Calc'] // 7

    def converter_ig_texto(sem):
        if sem <= 4: mes = "1º Mês"
        elif sem <= 8: mes = "2º Mês"
        elif sem <= 13: mes = "3º Mês"
        elif sem <= 17: mes = "4º Mês"
        elif sem <= 22: mes = "5º Mês"
        elif sem <= 27: mes = "6º Mês"
        elif sem <= 31: mes = "7º Mês"
        elif sem <= 35: mes = "8º Mês"
        elif sem <= 40: mes = "9º Mês"
        elif sem < 42: mes = "Pós-data"
        else: mes = "Pós-termo: ≥42 semanas"
        return f"{sem} semanas ({mes})"

    df['Idade gestacional atual'] = df['IG_Semanas'].apply(converter_ig_texto)

    # Captura flexível de últimas consultas e visitas
    col_ult_cons = buscar_coluna_flexivel(df, ['data da ultima consulta', 'ultima consulta'])
    col_ult_vis = buscar_coluna_flexivel(df, ['ultima visita', 'visitas domiciliares'])
    df['Data da última consulta'] = df[col_ult_cons] if col_ult_cons else "Não informada"
    df['Data da última visita'] = df[col_ult_vis] if col_ult_vis else "Não informada"

    def eval_c3(row):
        sem = row['IG_Semanas']
        ig_dias = row['IG_Dias_Calc']
        is_puerpera = ig_dias > 294
        
        def get_val_num(termos):
            c = buscar_coluna_flexivel(df, termos)
            if c and pd.notnull(row.get(c)):
                try:
                    nums = re.findall(r'\d+', str(row.get(c)))
                    if nums: return int(nums[0])
                except: pass
            return 0

        # Cálculo da meta esperada de consultas segundo a IG atual
        if sem <= 28: esperadas = max(1, sem // 4)
        elif sem <= 36: esperadas = 7 + max(0, (sem - 28) // 2)
        else: esperadas = 11 + max(0, sem - 36)
        esperadas_7 = min(esperadas, 7)
        esperadas_3 = min(max(1, sem // 10), 3)

        # Prática A (1ª Consulta <=12 sem)
        q_a = get_val_num(['12 semanas', 'atendimentos ate 12'])
        p_a = "Atendida" if q_a >= 1 else "Não atendida"

        # Prática B (Consultas de Pré-natal)
        q_b = min(get_val_num(['consultas de pre-natal', 'consultas pre natal']), 7)
        p_b = f"Atendida ({q_b}/{esperadas_7})" if q_b >= esperadas_7 else f"Não atendida ({q_b}/{esperadas_7})"

        # Prática C (Aferição de Pressão Arterial)
        q_c = min(get_val_num(['afericoes de pressao', 'pressao arterial']), 7)
        p_c = f"Atendida ({q_c}/{esperadas_7})" if q_c >= esperadas_7 else f"Não atendida ({q_c}/{esperadas_7})"

        # Prática D (Peso e Altura)
        q_d = min(get_val_num(['peso e altura', 'antropometria']), 7)
        p_d = f"Atendida ({q_d}/{esperadas_7})" if q_d >= esperadas_7 else f"Não atendida ({q_d}/{esperadas_7})"

        # Prática E (Visitas ACS Pré-natal)
        q_e = min(get_val_num(['visitas domiciliares', 'visitas acs']), 3)
        p_e = f"Atendida ({q_e}/{esperadas_3})" if q_e >= esperadas_3 else f"Não atendida ({q_e}/{esperadas_3})"

        # Prática F (dTpa >=20 sem)
        col_dtpa = buscar_coluna_flexivel(df, ['dtpa'])
        dtpa_v = str(row.get(col_dtpa, '') or '').strip() if col_dtpa else ''
        p_f = "Atendida" if dtpa_v and dtpa_v not in ['-', 'None', 'nan', ''] else ("Aguardando idade" if sem < 20 else "Não atendida")

        # Prática G (Exames 1º Trimestre)
        p_g = "Aguardando idade" if sem < 14 else "Não atendida"
        col_hiv1 = buscar_coluna_flexivel(df, ['hiv 1', 'hiv 1ºt'])
        if col_hiv1 and str(row.get(col_hiv1, '')).upper() in ['SIM', 'REALIZADO']: p_g = "Atendida"

        # Prática H (Exames 3º Trimestre)
        p_h = "Aguardando idade" if sem < 28 else "Não atendida"
        col_hiv3 = buscar_coluna_flexivel(df, ['hiv 3', 'hiv 3ºt'])
        if col_hiv3 and str(row.get(col_hiv3, '')).upper() in ['SIM', 'REALIZADO']: p_h = "Atendida"

        # Práticas I e J (Puerpério)
        col_cp = buscar_coluna_flexivel(df, ['consulta puerperio', 'puerperio consulta'])
        p_i = ("Atendida (1/1)" if col_cp and str(row.get(col_cp, '')).strip() not in ['-', '', 'None'] else "Não atendida (0/1)") if is_puerpera else "Não se aplica"
        
        col_vp = buscar_coluna_flexivel(df, ['visita puerperio', 'puerperio visita'])
        p_j = ("Atendida (1/1)" if col_vp and str(row.get(col_vp, '')).strip() not in ['-', '', 'None'] else "Não atendida (0/1)") if is_puerpera else "Não se aplica"

        # Prática K (Saúde Bucal)
        col_odonto = buscar_coluna_flexivel(df, ['odonto', 'odontologico'])
        p_k = "Atendida (1/1)" if col_odonto and get_val_num(['odonto']) >= 1 else "Não atendida (0/1)"

        praticas = [p_a, p_b, p_c, p_d, p_e, p_f, p_g, p_h, p_i, p_j, p_k]
        avaliaveis = [p for p in praticas if ("Atendida" in p or "Não atendida" in p) and "Não se aplica" not in p and "Aguardando" not in p]
        atendidas = [p for p in praticas if "Atendida" in p]
        score = (len(atendidas) / len(avaliaveis) * 100) if avaliaveis else 100.0

        return pd.Series([
            p_a, p_b, p_c, p_d, p_e, p_f, p_g, p_h, p_i, p_j, p_k, score
        ], index=[
            'Prática A — 1ª consulta até 12ª sem — status',
            'Prática B — consultas de pré-natal — status',
            'Prática C — aferição de pressão arterial — status',
            'Prática D — peso e altura — status',
            'Prática E — visitas ACS pré-natal — status',
            'Prática F — vacina dTpa a partir de 20 sem — status',
            'Prática G — exames 1º trimestre (HIV/Sífilis/HepB/HepC) — status',
            'Prática H — exames 3º trimestre (HIV/Sífilis) — status',
            'Prática I — consulta puerpério — status',
            'Prática J — visita ACS puerpério — status',
            'Prática K — atendimento odontológico — status',
            'Score_C3_%'
        ])

    res_c3 = df.apply(eval_c3, axis=1)
    for col in res_c3.columns: df[col] = res_c3[col]

    cols_nom = [
        'Microárea', 'Nome', 'CPF', 'Telefone celular', 'Endereço', 'Risco gestacional',
        'Idade gestacional atual', 'Data da última consulta', 'Data da última visita',
        'Prática A — 1ª consulta até 12ª sem — status',
        'Prática B — consultas de pré-natal — status',
        'Prática C — aferição de pressão arterial — status',
        'Prática D — peso e altura — status',
        'Prática E — visitas ACS pré-natal — status',
        'Prática F — vacina dTpa a partir de 20 sem — status',
        'Prática G — exames 1º trimestre (HIV/Sífilis/HepB/HepC) — status',
        'Prática H — exames 3º trimestre (HIV/Sífilis) — status',
        'Prática I — consulta puerpério — status',
        'Prática J — visita ACS puerpério — status',
        'Prática K — atendimento odontológico — status',
        'Score_C3_%'
    ]
    
    for c in cols_nom:
        if c not in df.columns: df[c] = "-"

    df_micro = df.groupby('Microárea').agg(
        Gestantes_Puerperas_Ativas=('Nome', 'count'),
        Score_Médio_C3=('Score_C3_%', 'mean'),
        Pct_100_Avaliável=('Score_C3_%', lambda x: (x == 100).mean() * 100)
    ).reset_index()

    return df_micro, df[cols_nom]

# --- OUTROS INDICADORES (MANTIDOS E INTEGRADOS) ---

def processar_c2(df, data_ref):
    df['Endereço'] = df.apply(montar_endereco, axis=1)
    col_nasc = buscar_coluna_flexivel(df, ['nascimento', 'dt_nasc'])
    if col_nasc:
        df['Nascimento'] = pd.to_datetime(df[col_nasc], format='%d/%m/%Y', errors='coerce')
        df['Idade_Dias'] = (data_ref - df['Nascimento']).dt.days
        df = df[df['Idade_Dias'] < 1096].copy()
    else: df['Idade_Dias'] = 300

    def eval_c2(row):
        dias = row.get('Idade_Dias', 300)
        col_cons1 = buscar_coluna_flexivel(df, ['primeira consulta', 'precoce'])
        p_a = "Não atendida (0/1)"
        if col_cons1 and row.get(col_cons1):
            try:
                dt = datetime.strptime(str(row.get(col_cons1))[:10], '%d/%m/%Y')
                if (dt - row.get('Nascimento', dt)).days <= 30: p_a = "Atendida (1/1)"
            except: pass
        if p_a != "Atendida (1/1)" and dias <= 30: p_a = "Aguardando idade"

        col_q_cons = buscar_coluna_flexivel(df, ['consultas', 'quantidade de consultas'])
        qtd_cons = min(int(pd.to_numeric(row.get(col_q_cons, 0), errors='coerce') or 0), 9) if col_q_cons else 0
        p_b = f"Atendida ({qtd_cons}/9)" if qtd_cons >= 9 else ("Aguardando idade" if dias < 730 else f"Não atendida ({qtd_cons}/9)")

        col_q_ant = buscar_coluna_flexivel(df, ['simultaneas', 'peso e altura', 'antropometria'])
        qtd_ant = min(int(pd.to_numeric(row.get(col_q_ant, 0), errors='coerce') or 0), 9) if col_q_ant else 0
        p_c = f"Atendida ({qtd_ant}/9)" if qtd_ant >= 9 else ("Aguardando idade" if dias < 730 else f"Não atendida ({qtd_ant}/9)")

        col_vis = buscar_coluna_flexivel(df, ['visita', 'visitas'])
        dts_v = extrair_datas_validas(row.get(col_vis), data_ref, 180) if col_vis else []
        p_d = "Atendida (2/2)" if len(dts_v) >= 2 else ("Aguardando idade" if dias < 180 else "Não atendida (0/2)")

        col_vac = buscar_coluna_flexivel(df, ['vacina', 'situacao vacinal'])
        p_e = "Atendida (1/1)" if col_vac and str(row.get(col_vac, '')).lower() in ['em dia', 'sim', 'completa'] else ("Aguardando idade" if dias < 365 else "Não atendida (0/1)")

        praticas = [p_a, p_b, p_c, p_d, p_e]
        avaliaveis = [p for p in praticas if "Atendida" in p or "Não atendida" in p]
        atendidas = [p for p in praticas if "Atendida" in p]
        score = (len(atendidas) / len(avaliaveis) * 100) if avaliaveis else 100.0

        return pd.Series([p_a, p_b, p_c, p_d, p_e, score],
                         index=['Prática A — 1ª consulta até 30d — status', 'Prática B — 9 consultas — status', 'Prática C — 9 antropometrias — status',
                                'Prática D — visitas ACS — status', 'Prática E — vacinação em dia — status', 'Score_C2_%'])

    res = df.apply(eval_c2, axis=1)
    for c in res.columns: df[c] = res[c]

    cols_nom = ['Microárea', 'Nome', 'CPF', 'Telefone celular', 'Endereço', 'Prática A — 1ª consulta até 30d — status', 'Prática B — 9 consultas — status', 
                'Prática C — 9 antropometrias — status', 'Prática D — visitas ACS — status', 'Prática E — vacinação em dia — status', 'Score_C2_%']
    for c in cols_nom:
        if c not in df.columns: df[c] = "-"

    df_micro = df.groupby('Microárea').agg(
        Crianças_Ativas=('Nome', 'count'),
        Score_Médio=('Score_C2_%', 'mean'),
        Pct_100_Avaliável=('Score_C2_%', lambda x: (x == 100).mean() * 100)
    ).reset_index()

    return df_micro, df[cols_nom]

def processar_c4(df, data_ref):
    df['Endereço'] = df.apply(montar_endereco, axis=1)

    def eval_c4(row):
        def check_janela(termos, dias):
            col = buscar_coluna_flexivel(df, termos)
            if col and row.get(col):
                try:
                    dt = datetime.strptime(str(row.get(col))[:10], '%d/%m/%Y')
                    if (data_ref - dt).days <= dias: return "Atendida (1/1)"
                except: pass
            return "Não atendida (0/1)"

        p_a = check_janela(['ultima consulta'], 183)
        p_b = check_janela(['pressao arterial'], 183)
        p_d = check_janela(['peso e altura'], 365)
        p_e = check_janela(['hemoglobina glicada', 'glicada'], 365)
        p_f = check_janela(['avaliacao dos pes', 'pes'], 365)

        col_vis = buscar_coluna_flexivel(df, ['visitas'])
        dts_v = extrair_datas_validas(row.get(col_vis), data_ref, 365) if col_vis else []
        p_c = "Não atendida (0/2)"
        if len(dts_v) >= 2:
            for i in range(len(dts_v)-1):
                if (dts_v[i+1] - dts_v[i]).days >= 30:
                    p_c = "Atendida (2/2)"
                    break

        pts = sum([1 for s in [p_a, p_b, p_c, p_d, p_e, p_f] if "Atendida" in s])
        score = (pts / 6.0) * 100
        return pd.Series([p_a, p_b, p_c, p_d, p_e, p_f, score],
                         index=['Prática A — consulta 6m — status', 'Prática B — pressão 6m — status', 'Prática C — 2 visitas 12m — status',
                                'Prática D — peso/altura 12m — status', 'Prática E — HbA1c 12m — status', 'Prática F — pés 12m — status', 'Score_C4_%'])

    res = df.apply(eval_c4, axis=1)
    for c in res.columns: df[c] = res[c]

    cols_nom = ['Microárea', 'Nome', 'CPF', 'Telefone celular', 'Endereço',
                'Prática A — consulta 6m — status', 'Prática B — pressão 6m — status', 'Prática C — 2 visitas 12m — status',
                'Prática D — peso/altura 12m — status', 'Prática E — HbA1c 12m — status', 'Prática F — pés 12m — status', 'Score_C4_%']
    for c in cols_nom:
        if c not in df.columns: df[c] = "-"

    df_micro = df.groupby('Microárea').agg(
        Diabéticos_Ativos=('Nome', 'count'),
        Score_Médio_C4=('Score_C4_%', 'mean'),
        Pct_100_Práticas=('Score_C4_%', lambda x: (x == 100).mean() * 100)
    ).reset_index()

    return df_micro, df[cols_nom]

def processar_c5(df, data_ref):
    df['Endereço'] = df.apply(montar_endereco, axis=1)

    def eval_c5(row):
        def check_janela(termos, dias):
            col = buscar_coluna_flexivel(df, termos)
            if col and row.get(col):
                try:
                    dt = datetime.strptime(str(row.get(col))[:10], '%d/%m/%Y')
                    if (data_ref - dt).days <= dias: return "Atendida (1/1)"
                except: pass
            return "Não atendida (0/1)"

        p_a = check_janela(['ultima consulta'], 183)
        p_b = check_janela(['pressao arterial'], 183)
        p_d = check_janela(['peso e altura'], 365)

        col_vis = buscar_coluna_flexivel(df, ['visitas'])
        dts_v = extrair_datas_validas(row.get(col_vis), data_ref, 365) if col_vis else []
        p_c = "Não atendida (0/2)"
        if len(dts_v) >= 2:
            for i in range(len(dts_v)-1):
                if (dts_v[i+1] - dts_v[i]).days >= 30:
                    p_c = "Atendida (2/2)"
                    break

        pts = sum([1 for s in [p_a, p_b, p_c, p_d] if "Atendida" in s])
        score = (pts / 4.0) * 100
        return pd.Series([p_a, p_b, p_c, p_d, score],
                         index=['Prática A — consulta 6m — status', 'Prática B — pressão 6m — status', 
                                'Prática C — 2 visitas 12m — status', 'Prática D — peso/altura 12m — status', 'Score_C5_%'])

    res = df.apply(eval_c5, axis=1)
    for c in res.columns: df[c] = res[c]

    cols_nom = ['Microárea', 'Nome', 'CPF', 'Telefone celular', 'Endereço',
                'Prática A — consulta 6m — status', 'Prática B — pressão 6m — status', 
                'Prática C — 2 visitas 12m — status', 'Prática D — peso/altura 12m — status', 'Score_C5_%']
    for c in cols_nom:
        if c not in df.columns: df[c] = "-"

    df_micro = df.groupby('Microárea').agg(
        Hipertensos_Ativos=('Nome', 'count'),
        Score_Médio_C5=('Score_C5_%', 'mean'),
        Pct_100_Práticas=('Score_C5_%', lambda x: (x == 100).mean() * 100)
    ).reset_index()

    return df_micro, df[cols_nom]

def processar_c6(df, data_ref):
    df['Endereço'] = df.apply(montar_endereco, axis=1)

    def eval_c6(row):
        col_med = buscar_coluna_flexivel(df, ['atendimento medico'])
        col_enf = buscar_coluna_flexivel(df, ['atendimento de enfermagem'])
        d_med = pd.to_numeric(row.get(col_med), errors='coerce') if col_med else None
        d_enf = pd.to_numeric(row.get(col_enf), errors='coerce') if col_enf else None
        p_a = "Atendida (1/1)" if ((pd.notnull(d_med) and 0 <= d_med <= 365) or (pd.notnull(d_enf) and 0 <= d_enf <= 365)) else "Não atendida (0/1)"

        col_ant = buscar_coluna_flexivel(df, ['peso e altura simultaneos', 'simultaneos'])
        qtd_ant = min(int(pd.to_numeric(row.get(col_ant, 0), errors='coerce') or 0), 2) if col_ant else 0
        p_b = f"Atendida ({qtd_ant}/2)" if qtd_ant >= 2 else f"Não atendida ({qtd_ant}/2)"

        col_vis = buscar_coluna_flexivel(df, ['visitas'])
        dts_v = extrair_datas_validas(row.get(col_vis), data_ref, 365) if col_vis else []
        p_c = "Não atendida (0/2)"
        if len(dts_v) >= 2:
            for i in range(len(dts_v)-1):
                if (dts_v[i+1] - dts_v[i]).days >= 30:
                    p_c = "Atendida (2/2)"
                    break

        col_flu = buscar_coluna_flexivel(df, ['influenza'])
        flu = str(row.get(col_flu, '') or '').strip() if col_flu else ''
        p_d = "Atendida (1/1)" if flu and flu not in ['-', 'Sem registro', 'None'] else "Não atendida (0/1)"

        pts = sum([1 for s in [p_a, p_b, p_c, p_d] if "Atendida" in s])
        score = (pts / 4.0) * 100
        return pd.Series([p_a, p_b, p_c, p_d, score],
                         index=['Prática A — consulta 12m — status', 'Prática B — 2 antropometrias 12m — status',
                                'Prática C — 2 visitas 12m — status', 'Prática D — vacina influenza 12m — status', 'Score_C6_%'])

    res = df.apply(eval_c6, axis=1)
    for c in res.columns: df[c] = res[c]

    cols_nom = ['Microárea', 'Nome', 'CPF', 'Telefone celular', 'Endereço',
                'Prática A — consulta 12m — status', 'Prática B — 2 antropometrias 12m — status',
                'Prática C — 2 visitas 12m — status', 'Prática D — vacina influenza 12m — status', 'Score_C6_%']
    for c in cols_nom:
        if c not in df.columns: df[c] = "-"

    df_micro = df.groupby('Microárea').agg(
        Idosos_Ativos=('Nome', 'count'),
        Score_Médio_C6=('Score_C6_%', 'mean'),
        Pct_100_Práticas=('Score_C6_%', lambda x: (x == 100).mean() * 100)
    ).reset_index()

    return df_micro, df[cols_nom]

def processar_c7(df, data_ref):
    df['Endereço'] = df.apply(montar_endereco, axis=1)
    col_nasc = buscar_coluna_flexivel(df, ['nascimento'])
    if col_nasc:
        df['Nascimento'] = pd.to_datetime(df[col_nasc], format='%d/%m/%Y', errors='coerce')
        df['Idade'] = ((data_ref - df['Nascimento']).dt.days / 365.25).fillna(30).astype(int)
    else: df['Idade'] = 30

    def eval_c7(row):
        idade = row['Idade']

        p_a = "Fora da faixa etária"
        if 25 <= idade <= 64:
            col_sol = buscar_coluna_flexivel(df, ['colo', 'colo de utero'])
            p_a = "Não atendida (0/1)"
            if col_sol and row.get(col_sol):
                try:
                    if (data_ref - datetime.strptime(str(row.get(col_sol))[:10], '%d/%m/%Y')).days <= 1095:
                        p_a = "Atendida (1/1)"
                except: pass

        p_b = "Fora da faixa etária"
        if 9 <= idade <= 14:
            col_hpv = buscar_coluna_flexivel(df, ['hpv'])
            hpv = str(row.get(col_hpv, '') or '').strip() if col_hpv else ''
            p_b = "Atendida (1/1)" if hpv and hpv not in ['-', 'Sem registro', 'None'] else "Não atendida (0/1)"

        p_c = "Fora da faixa etária"
        if 14 <= idade <= 69:
            col_ssr = buscar_coluna_flexivel(df, ['saude sexual', 'reprodutiva'])
            p_c = "Não atendida (0/1)"
            if col_ssr and row.get(col_ssr):
                try:
                    if (data_ref - datetime.strptime(str(row.get(col_ssr))[:10], '%d/%m/%Y')).days <= 365:
                        p_c = "Atendida (1/1)"
                except: pass

        p_d = "Fora da faixa etária"
        if 50 <= idade <= 69:
            col_mam = buscar_coluna_flexivel(df, ['mama', 'canser de mama'])
            p_d = "Não atendida (0/1)"
            if col_mam and row.get(col_mam):
                try:
                    if (data_ref - datetime.strptime(str(row.get(col_mam))[:10], '%d/%m/%Y')).days <= 730:
                        p_d = "Atendida (1/1)"
                except: pass

        praticas = [p_a, p_b, p_c, p_d]
        aplicaveis = [p for p in praticas if p != "Fora da faixa etária"]
        atendidas = [p for p in praticas if "Atendida (1/1)" in p]
        score = (len(atendidas) / len(aplicaveis) * 100) if aplicaveis else 100.0

        return pd.Series([p_a, p_b, p_c, p_d, score],
                         index=['Prática A — colo do útero 25-64a (36m) — status', 'Prática B — vacina HPV 9-14a — status',
                                'Prática C — saúde sexual/reprodutiva 14-69a (12m) — status', 'Prática D — câncer de mama 50-69a (24m) — status', 'Score_C7_%'])

    res = df.apply(eval_c7, axis=1)
    for c in res.columns: df[c] = res[c]

    cols_nom = ['Microárea', 'Nome', 'CPF', 'Idade', 'Telefone celular', 'Endereço',
                'Prática A — colo do útero 25-64a (36m) — status', 'Prática B — vacina HPV 9-14a — status',
                'Prática C — saúde sexual/reprodutiva 14-69a (12m) — status', 'Prática D — câncer de mama 50-69a (24m) — status', 'Score_C7_%']
    for c in cols_nom:
        if c not in df.columns: df[c] = "-"

    df_micro = df.groupby('Microárea').agg(
        Mulheres_Relatório=('Nome', 'count'),
        Score_Médio_C7=('Score_C7_%', 'mean'),
        Pct_100_Aplicáveis=('Score_C7_%', lambda x: (x == 100).mean() * 100)
    ).reset_index()

    return df_micro, df[cols_nom]

# --- INTERFACE GRÁFICA PRINCIPAL ---

mapa_indicadores = {
    "Indicador C2 — Puericultura (Crianças <3 anos)": ("C2", processar_c2, "INDICADOR C2 — PUERICULTURA", "Score_C2_%"),
    "Indicador C3 — Gestantes e Puérperas": ("C3", processar_c3, "INDICADOR C3 — SAÚDE DA GESTANTE E PUÉRPERA", "Score_C3_%"),
    "Indicador C4 — Cuidado da Pessoa com Diabetes": ("C4", processar_c4, "INDICADOR C4 — CUIDADO DA PESSOA COM DIABETES", "Score_C4_%"),
    "Indicador C5 — Cuidado da Pessoa com Hipertensão": ("C5", processar_c5, "INDICADOR C5 — CUIDADO DA PESSOA COM HIPERTENSÃO", "Score_C5_%"),
    "Indicador C6 — Cuidado da Pessoa Idosa": ("C6", processar_c6, "INDICADOR C6 — CUIDADO DA PESSOA IDOSA", "Score_C6_%"),
    "Indicador C7 — Prevenção do Câncer e Saúde da Mulher": ("C7", processar_c7, "INDICADOR C7 — SAÚDE DA MULHER E PREVENÇÃO DO CÂNCER", "Score_C7_%")
}

opção_selecionada = st.selectbox(
    "Selecione o Indicador para Auditoria:",
    list(mapa_indicadores.keys())
)

uploaded_file = st.file_uploader("Anexe o relatório CSV exportado do e-SUS PEC:", type=["csv"])

if uploaded_file and st.button("🚀 Processar Dashboard Agora"):
    with st.spinner("Lendo relatório e aplicando regras oficiais da Nota Metodológica..."):
        bytes_data = uploaded_file.getvalue()
        df_bruto, data_ref = carregar_e_limpar_csv(bytes_data)
        
        sigla, func_proc, titulo_doc, col_score = mapa_indicadores[opção_selecionada]
        
        df_micro, df_nom = func_proc(df_bruto, data_ref)
        
        cols_status = [c for c in df_nom.columns if "Prática" in c or "status" in c]
        excel_bytes = gerar_excel_padrao(df_micro, df_nom, titulo_doc, data_ref, cols_status, col_score)
        
        st.success(f"✅ Dashboard do {sigla} gerado com sucesso!")
        st.download_button(
            label=f"📥 Baixar Dashboard_Indicador_{sigla}.xlsx",
            data=excel_bytes,
            file_name=f"Dashboard_Indicador_{sigla}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
